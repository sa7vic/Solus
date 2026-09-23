"""
Real Word document (.docx) generation using python-docx.
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

def write_document(out_path: str, title: str, subtitle: str, sections: list[dict], citations: list[dict]) -> str:
    doc = Document()
    
    # Title
    h1 = doc.add_heading(title or "Document", level=0)
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    if subtitle:
        p_sub = doc.add_paragraph(subtitle)
        p_sub.runs[0].font.size = Pt(12)
        p_sub.runs[0].font.italic = True
        p_sub.runs[0].font.color.rgb = RGBColor(100, 110, 120)

    for sec in sections:
        doc.add_heading(sec.get("heading", "Section"), level=1)
        for bullet in sec.get("bullets", []):
            text = bullet.get("text", "")
            cit = bullet.get("citation")
            if cit:
                text = f"{text} [{cit}]"
            p = doc.add_paragraph(text, style="List Bullet")
            p.runs[0].font.size = Pt(10.5)

    if citations:
        doc.add_heading("References & Grounded Citations", level=1)
        for c in citations:
            p = doc.add_paragraph(f"[{c['n']}] {c['source']}")
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(80, 80, 80)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path

